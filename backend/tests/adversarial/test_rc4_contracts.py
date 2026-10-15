"""RC4: adversarial OOXML intake, retention and recovery regression gates."""
from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import main
from app.security.intake import IntakeError, validate_pptx
from app.storage.sessions import SessionStore

REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
PRES = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
VENUE = {"active_image_height_m": 1.8, "farthest_viewer_distance_m": 10.8, "measurement_basis": "MEASURED"}
CONSTRAINTS = {"max_mutations": 2, "max_scale_factor": 1.6,
               "minimum_target_coverage": 0.75, "protect_titles": True,
               "semantic_rewrite": False}


def _valid(tmp_path):
    prs = Presentation()
    for i in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000000, 800000).text = f"UNIQUE SLIDE {i + 1}"
    out = tmp_path / "valid.pptx"
    prs.save(out)
    return out


def _alter(src, target, part, edit):
    with zipfile.ZipFile(src) as before, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as after:
        for info in before.infolist():
            data = before.read(info.filename)
            if info.filename == part:
                root = ET.fromstring(data)
                edit(root)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            after.writestr(info, data)


@pytest.mark.parametrize("kind", ["duplicate_identical", "duplicate_conflicting", "bad_root", "missing_id", "bad_mode", "foreign_element"])
def test_opc_relationship_ambiguity_rejected_before_allocation(tmp_path, kind):
    source = _valid(tmp_path)
    rels = "ppt/_rels/presentation.xml.rels"
    out = tmp_path / f"{kind}.pptx"
    def edit(root):
        if kind.startswith("duplicate"):
            first = next(e for e in root if e.get("Type", "").endswith("/slide"))
            dup = ET.SubElement(root, first.tag, dict(first.attrib))
            if kind == "duplicate_conflicting":
                dup.set("Target", "slides/slide2.xml")
        elif kind == "bad_root":
            root.tag = "Relationships"
        elif kind == "missing_id":
            next(iter(root)).attrib.pop("Id")
        elif kind == "bad_mode":
            next(iter(root)).set("TargetMode", "Maybe")
        elif kind == "foreign_element":
            ET.SubElement(root, "{http://example.invalid}Relationship", {"Id":"rId9"})
    _alter(source, out, rels, edit)
    with pytest.raises(IntakeError):
        validate_pptx(out)
    with TestClient(main.app) as client:
        with out.open("rb") as file:
            response = client.post("/api/upload", files={"file": (out.name, file, "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert response.status_code == 400


def test_empty_logical_slide_list_rejected(tmp_path):
    source = _valid(tmp_path)
    target = tmp_path / "orphan.pptx"
    def edit(root):
        node = root.find(PRES + "sldIdLst")
        assert node is not None
        node.clear()
    _alter(source, target, "ppt/presentation.xml", edit)
    with pytest.raises(IntakeError, match="Logical slide count"):
        validate_pptx(target)


def test_duplicate_logical_slides_rejected(tmp_path):
    source = _valid(tmp_path)
    target = tmp_path / "reusedslide.pptx"
    def edit(root):
        node = root.find(PRES + "sldIdLst")
        node[1].set("id", node[0].get("id"))
    _alter(source, target, "ppt/presentation.xml", edit)
    with pytest.raises(IntakeError, match="Duplicate or missing logical slide"):
        validate_pptx(target)


def test_normal_slide_reorder_keeps_logical_inventory(tmp_path):
    source = _valid(tmp_path)
    out = tmp_path / "reordered.pptx"
    def edit(root):
        node = root.find(PRES + "sldIdLst")
        a, b = list(node)
        node.remove(a)
        node.append(a)
    _alter(source, out, "ppt/presentation.xml", edit)
    assert validate_pptx(out).slide_count == 2
    assert [s.shapes[0].text for s in Presentation(out).slides] == ["UNIQUE SLIDE 2", "UNIQUE SLIDE 1"]


def test_first_admission_after_expiry_succeeds(tmp_path):
    class Clock:
        now = 1234.0
        def __call__(self): return self.now
    clock = Clock()
    store = SessionStore(tmp_path / "sessions", clock=clock, max_sessions=1, per_owner_sessions=1,
                         ttl_seconds=600, provisional_ttl_seconds=60)
    try:
        old = store.create("one.pptx", b"one", "a", owner_key="person")
        clock.now += 100
        assert store.begin_intake("person") is None
        store.finish_intake("person")
        assert not old.source_path.exists()
    finally:
        store.stop_reaper()


@pytest.mark.resource
def test_completed_replan_refused_keeps_verified_artifact_and_invalidates_old_approval(tmp_path, monkeypatch):
    store = SessionStore(tmp_path / "sessions")
    monkeypatch.setattr(main, "STORE", store)
    try:
        with TestClient(main.app) as client:
            created = client.post("/api/demo").json()
            sid = created["session_id"]
            headers = {"X-Venue-Token": created["session_token"]}
            uri = f"/api/sessions/{sid}"
            assert client.post(uri + "/analyze", headers=headers, json={"venue": VENUE}).status_code == 200
            plan = client.post(uri + "/plan", headers=headers, json={"constraints": CONSTRAINTS}).json()
            assert plan["state"] == "AWAITING_APPROVAL", plan
            approved = client.post(uri + "/approve-and-apply", headers=headers, json={
                "approve": True, "plan_hash": plan["plan_hash"],
                "candidate_ids": plan["decision"]["selected_candidate_ids"],
                "idempotency_key": "rc4-initial-approval",
            })
            assert approved.status_code == 200 and approved.json()["state"] == "VERIFIED", approved.text
            original = client.get(uri + "/output", headers=headers)
            assert original.status_code == 200
            def must_not_choose(*args, **kwargs):
                raise AssertionError("Completed-session replanning must never invoke planner")
            monkeypatch.setattr(main.PLANNER, "choose", must_not_choose)
            second = client.post(uri + "/plan", headers=headers, json={"constraints": CONSTRAINTS})
            assert second.status_code == 409, second.text
            assert "new session" in second.text.lower()
            status = client.get(uri, headers=headers).json()
            assert status["state"] == "VERIFIED" and status["has_output"] is True
            assert status["pending_plan"] is None
            assert client.get(uri + "/output", headers=headers).content == original.content
            replay = client.post(uri + "/approve-and-apply", headers=headers, json={
                "approve": True, "plan_hash": plan["plan_hash"],
                "candidate_ids": plan["decision"]["selected_candidate_ids"],
                "idempotency_key": "rc4-replay",
            })
            assert replay.status_code == 409
    finally:
        store.stop_reaper()


@pytest.mark.resource
def test_pending_plan_recovery_preserves_exact_planner_decision(tmp_path, monkeypatch):
    store = SessionStore(tmp_path / "sessions")
    monkeypatch.setattr(main, "STORE", store)
    try:
        with TestClient(main.app) as client:
            created = client.post("/api/demo").json()
            uri = f"/api/sessions/{created['session_id']}"
            headers = {"X-Venue-Token": created["session_token"]}
            assert client.post(uri + "/analyze", headers=headers, json={"venue": VENUE}).status_code == 200
            planned = client.post(uri + "/plan", headers=headers, json={"constraints": CONSTRAINTS}).json()
            assert planned["state"] == "AWAITING_APPROVAL", planned
            recovered = client.get(uri, headers=headers).json()["pending_plan"]
            assert recovered is not None
            assert recovered["decision"] == planned["decision"]
            assert recovered["plan_hash"] == planned["plan_hash"]
    finally:
        store.stop_reaper()


def test_live_shared_session_root_is_rejected_without_deleting_data(tmp_path):
    root = tmp_path / "shared"
    first = SessionStore(root)
    session = first.create("a.pptx", b"owned-content", "digest", owner_key="a")
    try:
        with pytest.raises(RuntimeError, match="session_root_already_in_use"):
            SessionStore(root)
        assert session.source_path.read_bytes() == b"owned-content"
    finally:
        first.stop_reaper()
    second = SessionStore(root)
    try:
        assert not session.source_path.exists()  # Exclusive clean-start recovery only.
    finally:
        second.stop_reaper()


def test_readiness_and_frontend_have_consistent_operational_headers(monkeypatch):
    with TestClient(main.app) as client:
        homepage = client.get("/")
        assert homepage.status_code == 200
        assert "no-cache" in homepage.headers.get("cache-control", "")
        monkeypatch.setattr(main, "libreoffice_version", lambda: None)
        readiness = client.get("/health/readiness")
        assert readiness.status_code == 503
        assert readiness.json()["status"] == "degraded"


def test_font_preservation_requires_every_explicit_run_to_receive_approved_scale(tmp_path):
    from pptx.util import Pt
    from app.pptx.mutations import ScaleMutation, apply_scale_plan
    from app.pptx.package_integrity import package_integrity_reasons
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, 4000000, 1000000)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.clear()
    for value in ("First", " second"):
        run = paragraph.add_run()
        run.text = value
        run.font.size = Pt(16)
    source = tmp_path / "runs.pptx"
    prs.save(source)
    output = tmp_path / "scaled.pptx"
    apply_scale_plan(source, output, [ScaleMutation(0, shape.shape_id, 1.25)])
    expected = {(0, shape.shape_id): 1.25}
    assert package_integrity_reasons(source, output, {(0, shape.shape_id)}, expected) == []
    broken = tmp_path / "partial-scale.pptx"
    drawing = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    def undo_one_run(root):
        props = root.findall(f".//{drawing}rPr")
        assert len(props) >= 2
        props[0].set("sz", "1600")
    _alter(output, broken, "ppt/slides/slide1.xml", undo_one_run)
    assert "unapproved_slide_appearance_or_structure_changed" in package_integrity_reasons(
        source, broken, {(0, shape.shape_id)}, expected)

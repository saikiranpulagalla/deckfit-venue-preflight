"""RC3 independent regressions for the cross-audit confirmed release blockers."""
from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import main
from app.pptx.mutations import ScaleMutation, apply_scale_plan
from app.pptx.package_integrity import package_integrity_reasons
from app.storage.sessions import SessionStore

VENUE = {'active_image_height_m':1.8,'farthest_viewer_distance_m':10.8,'measurement_basis':'MEASURED'}
LIMITS = {'max_mutations':2,'max_scale_factor':1.6,'minimum_target_coverage':.75,
          'protect_titles':True,'semantic_rewrite':False}


def _make_changed(src:Path,dst:Path,edit):
    with zipfile.ZipFile(src) as z1,zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as z2:
        for info in z1.infolist():
            b=z1.read(info.filename)
            if info.filename=='ppt/slides/slide1.xml':
                root=ET.fromstring(b)
                edit(root)
                b=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            z2.writestr(info,b)


def _sample_font_only(tmp_path):
    src=Path(__file__).resolve().parents[3]/'demo/venue-agent-demo.pptx'
    prs=Presentation(src)
    first=next((i,s.shape_id) for i,slide in enumerate(prs.slides) for s in slide.shapes
               if s.has_text_frame and any(r.text.strip() and r.font.size is not None
                                           for p in s.text_frame.paragraphs for r in p.runs))
    out=tmp_path/'repaired.pptx'
    apply_scale_plan(src,out,[ScaleMutation(first[0],first[1],1.15)])
    return src,out,first


def test_font_only_allows_only_approved_shapes_and_keeps_whole_package(tmp_path):
    src,out,first=_sample_font_only(tmp_path)
    assert package_integrity_reasons(src,out,{first}) == []
    assert 'unapproved_slide_appearance_or_structure_changed' in package_integrity_reasons(src,out,set())


def test_background_and_text_color_are_protected_against_white_on_white(tmp_path):
    src,out,first=_sample_font_only(tmp_path)
    P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
    A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
    def background(root):
        common=root.find(f'{P}cSld')
        bg=ET.SubElement(common,f'{P}bg')
        bgpr=ET.SubElement(bg,f'{P}bgPr')
        solid=ET.SubElement(bgpr,f'{A}solidFill')
        ET.SubElement(solid,f'{A}srgbClr',{'val':'000000'})
    altered=tmp_path/'background.pptx'
    _make_changed(out,altered,background)
    assert 'unapproved_slide_appearance_or_structure_changed' in package_integrity_reasons(src,altered,{first})
    def color(root):
        pr=root.find(f'.//{A}rPr')
        assert pr is not None
        fill=ET.SubElement(pr,f'{A}solidFill')
        ET.SubElement(fill,f'{A}srgbClr',{'val':'FFFFFF'})
    altered2=tmp_path/'color.pptx'
    _make_changed(out,altered2,color)
    assert 'unapproved_slide_appearance_or_structure_changed' in package_integrity_reasons(src,altered2,{first})


def test_other_parts_are_never_mutable(tmp_path):
    src,out,first=_sample_font_only(tmp_path)
    name='ppt/presentation.xml'
    altered=tmp_path/'changed.pptx'
    with zipfile.ZipFile(out) as a,zipfile.ZipFile(altered,'w',zipfile.ZIP_DEFLATED) as b:
        for item in a.infolist():
            data=a.read(item.filename)
            if item.filename==name:
                data=data.replace(b'</p:presentation>',b'<p:extLst/></p:presentation>')
            b.writestr(item,data)
    assert package_integrity_reasons(src,altered,{first})


def test_unclaimed_session_expires_faster_than_authenticated_session(tmp_path):
    class Clock:
        now=2000.0
        def __call__(self):return self.now
    clock=Clock()
    store=SessionStore(tmp_path/'sessions',clock=clock,ttl_seconds=2700,provisional_ttl_seconds=120)
    try:
        a=store.create('a.pptx',b'a',hashlib.sha256(b'a').hexdigest(),owner_key='a')
        token=a.take_issued_capability_token()
        assert a.expires_at==2120
        assert store.authorize(a.session_id,token)
        assert a.expires_at==4700
        b=store.create('b.pptx',b'b',hashlib.sha256(b'b').hexdigest(),owner_key='a')
        clock.now=2121
        assert store.get(a.session_id)
        with pytest.raises(KeyError):store.get(b.session_id)
    finally:store.stop_reaper()


def test_tampered_output_status_does_not_advertise_download(tmp_path,monkeypatch):
    store=SessionStore(tmp_path/'sessions')
    monkeypatch.setattr(main,'STORE',store)
    try:
        s=store.create('a.pptx',b'a',hashlib.sha256(b'a').hexdigest(),owner_key='a')
        token=s.take_issued_capability_token()
        s.output_path=s.root/'output.pptx'
        s.output_path.write_bytes(b'original-result')
        s.output_sha256=hashlib.sha256(b'original-result').hexdigest()
        s.state=main.WorkflowState.VERIFIED
        s.output_path.write_bytes(b'tampered')
        with TestClient(main.app) as client:
            res=client.get(f'/api/sessions/{s.session_id}',headers={'X-Venue-Token':token})
        assert res.status_code==200
        assert res.json()['has_output'] is False
        assert res.json()['state']=='FAILED'
    finally:store.stop_reaper()


@pytest.mark.resource
def test_current_plan_hash_is_unique_across_replan_and_reanalysis(tmp_path,monkeypatch):
    store=SessionStore(tmp_path/'sessions')
    monkeypatch.setattr(main,'STORE',store)
    try:
        with TestClient(main.app) as c:
            item=c.post('/api/demo').json();sid=item['session_id'];h={'X-Venue-Token':item['session_token']}
            uri=f'/api/sessions/{sid}'
            a=c.post(uri+'/analyze',headers=h,json={'venue':VENUE});assert a.status_code==200,a.text
            p1=c.post(uri+'/plan',headers=h,json={'constraints':LIMITS});assert p1.status_code==200,p1.text
            assert p1.json()['state']=='AWAITING_APPROVAL',p1.text
            p2=c.post(uri+'/plan',headers=h,json={'constraints':LIMITS});assert p2.status_code==200,p2.text
            assert p1.json()['plan_hash']!=p2.json()['plan_hash']
            old=c.post(uri+'/approve-and-apply',headers=h,json={
                'candidate_ids':p1.json()['decision']['selected_candidate_ids'],
                'plan_hash':p1.json()['plan_hash'],'idempotency_key':'old-plan','approve':True})
            assert old.status_code==409,old.text
            status=c.get(uri,headers=h);assert status.json()['pending_plan']['venue']==VENUE
            assert status.json()['pending_plan']['constraints']['max_mutations']==2
            a=c.post(uri+'/analyze',headers=h,json={'venue':VENUE});assert a.status_code==200,a.text
            p3=c.post(uri+'/plan',headers=h,json={'constraints':LIMITS});assert p3.status_code==200,p3.text
            assert p2.json()['plan_hash']!=p3.json()['plan_hash']
            stale=c.post(uri+'/approve-and-apply',headers=h,json={
                'candidate_ids':p2.json()['decision']['selected_candidate_ids'],
                'plan_hash':p2.json()['plan_hash'],'idempotency_key':'old-analysis','approve':True})
            assert stale.status_code==409,stale.text
    finally:store.stop_reaper()


@pytest.mark.resource
def test_completed_output_prevents_reanalysis_and_retains_last_usable_output(tmp_path,monkeypatch):
    store=SessionStore(tmp_path/'sessions')
    monkeypatch.setattr(main,'STORE',store)
    with TestClient(main.app) as c:
        item=c.post('/api/demo').json();sid=item['session_id'];h={'X-Venue-Token':item['session_token']}
        uri=f'/api/sessions/{sid}'
        a=c.post(uri+'/analyze',headers=h,json={'venue':VENUE});assert a.status_code==200,a.text
        p=c.post(uri+'/plan',headers=h,json={'constraints':LIMITS});assert p.json()['state']=='AWAITING_APPROVAL',p.text
        approved=c.post(uri+'/approve-and-apply',headers=h,json={
            'candidate_ids':p.json()['decision']['selected_candidate_ids'],
            'plan_hash':p.json()['plan_hash'],'idempotency_key':'first','approve':True})
        assert approved.status_code==200,approved.text
        assert approved.json()['state']=='VERIFIED',approved.text
        first=c.get(uri+'/output',headers=h);assert first.status_code==200
        def should_never_run(*args, **kwargs):
            raise AssertionError('Completed session reanalysis must not invoke the analyzer')
        monkeypatch.setattr(main.ANALYZER, 'analyze', should_never_run)
        refused=c.post(uri+'/analyze',headers=h,json={'venue':{**VENUE,'farthest_viewer_distance_m':11.0}})
        assert refused.status_code==409,refused.text
        assert 'new session' in refused.text.lower()
        status=c.get(uri,headers=h);assert status.json()['has_output'] is True
        assert status.json()['state']=='VERIFIED'
        assert status.json()['analysis_venue']==VENUE
        assert c.get(uri+'/output',headers=h).content==first.content
        assert status.json()['pending_plan'] is None
    store.stop_reaper()


def test_nested_svg_external_url_is_refused(tmp_path):
    from app.security.intake import IntakeError, validate_pptx
    src, _out, _target = _sample_font_only(tmp_path)
    altered = tmp_path/'external-svg.pptx'
    types = '[Content_Types].xml'
    ns = '{http://schemas.openxmlformats.org/package/2006/content-types}'
    with zipfile.ZipFile(src) as origin, zipfile.ZipFile(altered,'w',zipfile.ZIP_DEFLATED) as result:
        for entry in origin.infolist():
            data=origin.read(entry.filename)
            if entry.filename==types:
                root=ET.fromstring(data)
                ET.SubElement(root,ns+'Default',{'Extension':'svg','ContentType':'image/svg+xml'})
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            result.writestr(entry,data)
        result.writestr('ppt/media/image99.svg',
                        b'<svg xmlns="http://www.w3.org/2000/svg"><image href="https://example.invalid/track"/></svg>')
    with pytest.raises(IntakeError,match='unsupported external reference'):
        validate_pptx(altered)


def test_readiness_version_lookup_cached(monkeypatch):
    from app.pptx import renderer
    from types import SimpleNamespace
    calls=[]
    renderer.libreoffice_version.cache_clear()
    monkeypatch.setattr(renderer.shutil,'which',lambda exe: '/opt/example/soffice')
    monkeypatch.setattr(renderer.subprocess,'run',lambda *a,**kw:(calls.append(a) or SimpleNamespace(stdout='LibreOffice patched',stderr='')))
    try:
        assert renderer.libreoffice_version()=='LibreOffice patched'
        assert renderer.libreoffice_version()=='LibreOffice patched'
        assert len(calls)==1
    finally:
        renderer.libreoffice_version.cache_clear()


@pytest.mark.resource
def test_real_verifier_rejects_unapproved_visual_property_change(tmp_path):
    from app.repair.verifier import verify_output
    from app.domain.models import WorkflowState
    src=Path(__file__).resolve().parents[3]/'backend/tests/fixtures/small_text.pptx'
    analyzer=main.ANALYZER
    v=main.VenueProfile(**VENUE)
    before=analyzer.analyze(src,v,tmp_path/'before')
    target=next(r for r in before.results if r.element.shape_id is not None
                and r.element.role=='BODY' and r.state.value=='BELOW_TARGET')
    original=tmp_path/'legit.pptx'; altered=tmp_path/'changed-background.pptx'
    scale=1.3
    apply_scale_plan(src,original,[ScaleMutation(target.element.slide_index,target.element.shape_id,scale)])
    A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
    P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
    def tamper_background(root):
        common=root.find(f'{P}cSld')
        bg=ET.SubElement(common,f'{P}bg'); bgpr=ET.SubElement(bg,f'{P}bgPr')
        ET.SubElement(ET.SubElement(bgpr,f'{A}solidFill'),f'{A}srgbClr',{'val':'000000'})
    _make_changed(original,altered,tamper_background)
    report=verify_output(src,altered,before,[target.element.element_id],analyzer,tmp_path/'verification',
                         expected_scales={(target.element.slide_index,target.element.shape_id):scale})
    assert report.final_state!=WorkflowState.VERIFIED
    assert 'unapproved_slide_appearance_or_structure_changed' in report.reasons

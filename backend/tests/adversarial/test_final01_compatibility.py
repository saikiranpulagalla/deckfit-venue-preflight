"""FINAL-01: compatibility wrappers must not escape visual coverage at intake.

API controls use actual local LibreOffice; no analyzer/verifier/renderer mocks.
"""
from __future__ import annotations

import copy
import hashlib
import posixpath
from pathlib import Path
from zipfile import ZipFile

from fastapi.testclient import TestClient
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches
import pytest

from app import main
from app.security.intake import IntakeError, validate_pptx
from app.storage.sessions import SessionStore

MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def _fixture(source: Path, output: Path, *, scope='slide', content='visible',
             prefix='mc', namespace=MC, hidden=False, outside_tree=False,
             branches='both') -> Path:
    prs = Presentation(source)
    slide = prs.slides[0]
    part = {'slide': slide.part, 'layout': slide.slide_layout.part,
            'master': slide.slide_layout.slide_master.part}[scope]
    name = str(part.partname).lstrip('/')
    rectangle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        prs.slide_width-Inches(1.2), prs.slide_height-Inches(1.2), Inches(.7), Inches(.7))
    rectangle.fill.solid()
    rectangle.fill.fore_color.rgb = RGBColor(255, 0, 0)
    rectangle.line.fill.background()
    visual = rectangle._element
    with ZipFile(source) as src, ZipFile(output, 'w') as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == name:
                root = etree.fromstring(data)
                tree = root.find('.//{' + P + '}spTree')
                if hidden and scope == 'slide':
                    root.set('show', '0')
                if content == 'direct':
                    tree.append(copy.deepcopy(visual))
                else:
                    wrapper = etree.Element('{' + namespace + '}AlternateContent', nsmap={
                        prefix: namespace, 'a14': 'http://schemas.microsoft.com/office/drawing/2010/main'})
                    choice = etree.SubElement(wrapper, '{' + namespace + '}Choice', Requires='a14')
                    fallback = etree.SubElement(wrapper, '{' + namespace + '}Fallback')
                    for branch, label in ((choice, 'choice'), (fallback, 'fallback')):
                        if branches not in {'both', label}:
                            continue
                        if content == 'visible':
                            branch.append(copy.deepcopy(visual))
                        elif content == 'metadata':
                            etree.SubElement(branch, '{urn:deckfit:test:metadata}metadata').text = 'inert'
                    if outside_tree:
                        extensions = etree.SubElement(root, '{' + P + '}extLst')
                        etree.SubElement(extensions, '{' + P + '}ext', uri='urn:deckfit:test').append(wrapper)
                    else:
                        tree.append(wrapper)
                data = etree.tostring(root, encoding='utf-8', xml_declaration=True)
            dst.writestr(info, data)
    return output


@pytest.mark.parametrize('scope', ['slide', 'layout', 'master'])
@pytest.mark.parametrize('content', ['visible', 'empty', 'metadata'])
def test_compatibility_content_refused_in_visual_parts(project_root, tmp_path, scope, content):
    path = _fixture(project_root/'demo/venue-agent-demo.pptx', tmp_path/'wrapper.pptx',
                    scope=scope, content=content)
    with pytest.raises(IntakeError, match=r'Unsupported presentation compatibility content'):
        validate_pptx(path)


@pytest.mark.parametrize('branches', ['choice', 'fallback', 'both'])
def test_branch_presence_and_duplicate_visuals_never_bypass_refusal(project_root, tmp_path, branches):
    path = _fixture(project_root/'demo/venue-agent-demo.pptx', tmp_path/'branches.pptx', branches=branches)
    with pytest.raises(IntakeError, match='mc:AlternateContent'):
        validate_pptx(path)


def test_namespace_not_literal_prefix_controls_refusal(project_root, tmp_path):
    source = project_root/'demo/venue-agent-demo.pptx'
    alias = _fixture(source, tmp_path/'alias.pptx', prefix='compat')
    with pytest.raises(IntakeError, match='mc:AlternateContent'):
        validate_pptx(alias)
    # Same local name in an unrelated namespace is not the compatibility format.
    metadata = _fixture(source, tmp_path/'unrelated.pptx', content='metadata',
                        namespace='urn:deckfit:test:not-mc', outside_tree=True)
    validate_pptx(metadata)


@pytest.mark.parametrize('scope', ['slide', 'layout', 'master'])
def test_metadata_wrapper_outside_shape_tree_is_deliberately_refused(project_root, tmp_path, scope):
    path = _fixture(project_root/'demo/venue-agent-demo.pptx', tmp_path/'metadata.pptx',
                    scope=scope, content='metadata', outside_tree=True)
    with pytest.raises(IntakeError, match='mc:AlternateContent'):
        validate_pptx(path)


def test_hidden_slide_does_not_exempt_compatibility_markup(project_root, tmp_path):
    source = project_root/'demo/venue-agent-demo.pptx'
    hidden = _fixture(source, tmp_path/'hidden-wrapper.pptx', hidden=True)
    with pytest.raises(IntakeError, match='mc:AlternateContent'):
        validate_pptx(hidden)
    validate_pptx(_fixture(source, tmp_path/'hidden-direct.pptx', hidden=True, content='direct'))


@pytest.mark.parametrize('scope', ['layout', 'master'])
def test_content_type_identifies_relocated_visual_parts(project_root, tmp_path, scope):
    source = project_root/'demo/venue-agent-demo.pptx'
    prs = Presentation(source)
    part = prs.slides[0].slide_layout.part if scope == 'layout' else prs.slides[0].slide_layout.slide_master.part
    old = str(part.partname).lstrip('/')
    new = 'ppt/custom/' + posixpath.basename(old)
    wrapped = _fixture(source, tmp_path/'wrapped.pptx', scope=scope)
    def relname(name):
        directory, leaf = posixpath.split(name)
        return posixpath.join(directory, '_rels', leaf + '.rels')
    moved = {old:new, relname(old):relname(new)}
    target = tmp_path/'relocated.pptx'
    with ZipFile(wrapped) as src, ZipFile(target, 'w') as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == '[Content_Types].xml':
                root = etree.fromstring(data)
                for element in root:
                    if element.get('PartName') == '/' + old:
                        element.set('PartName', '/' + new)
                data = etree.tostring(root, encoding='utf-8', xml_declaration=True)
            elif info.filename.endswith('.rels'):
                root = etree.fromstring(data)
                directory = posixpath.dirname(info.filename).removesuffix('/_rels')
                for element in root:
                    if element.get('TargetMode', 'Internal') == 'Internal':
                        resolved = posixpath.normpath(posixpath.join(directory, element.get('Target', '')))
                        if resolved == old:
                            element.set('Target', posixpath.relpath(new, directory))
                data = etree.tostring(root, encoding='utf-8', xml_declaration=True)
            dst.writestr(moved.get(info.filename, info.filename), data)
    # The relocated part is legal and reachable, not a dangling-relationship case.
    loaded = Presentation(target)
    actual = loaded.slides[0].slide_layout.part if scope == 'layout' else loaded.slides[0].slide_layout.slide_master.part
    assert str(actual.partname).lstrip('/') == new
    with pytest.raises(IntakeError, match='mc:AlternateContent'):
        validate_pptx(target)
    control = tmp_path/'relocated-without-wrapper.pptx'
    with ZipFile(target) as src, ZipFile(control, 'w') as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == new:
                root = etree.fromstring(data)
                for wrapper in list(root.iter('{' + MC + '}AlternateContent')):
                    wrapper.getparent().remove(wrapper)
                data = etree.tostring(root, encoding='utf-8', xml_declaration=True)
            dst.writestr(info, data)
    validate_pptx(control)


@pytest.mark.parametrize('scope', ['slide', 'layout', 'master'])
def test_real_api_rejects_wrapper_without_creating_session(project_root, tmp_path, monkeypatch, scope):
    store = SessionStore(tmp_path/'sessions')
    monkeypatch.setattr(main, 'STORE', store)
    path = _fixture(project_root/'demo/venue-agent-demo.pptx', tmp_path/'upload.pptx', scope=scope)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    with TestClient(main.app) as client:
        response = client.post('/api/upload', files={'file': ('upload.pptx', path.read_bytes(),
                               'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
        assert response.status_code == 400, response.text
        assert 'mc:AlternateContent' in response.json()['detail']
        assert not store._items and not store._pending_intakes and not store._pending_creates
        assert not list(store.base.glob('session-*'))
        assert 'session_token' not in response.json()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


@pytest.mark.resource
@pytest.mark.parametrize('case', ['direct', 'small_text', 'demo'])
def test_real_api_positive_and_direct_control(project_root, tmp_path, monkeypatch, case):
    store = SessionStore(tmp_path/'sessions')
    monkeypatch.setattr(main, 'STORE', store)
    # Force only the supported local deterministic mode, never a paid API call.
    from app.agent.planner import DeterministicPlanner
    monkeypatch.setattr(main, 'PLANNER', DeterministicPlanner())
    source = project_root/('backend/tests/fixtures/small_text.pptx' if case == 'small_text'
                           else 'demo/venue-agent-demo.pptx')
    if case == 'direct':
        source = _fixture(source, tmp_path/'direct.pptx', content='direct')
    original = hashlib.sha256(source.read_bytes()).hexdigest()
    with TestClient(main.app) as client:
        response = client.post('/api/upload', files={'file': ('input.pptx', source.read_bytes(),
                               'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
        assert response.status_code == 200, response.text
        created = response.json()
        base = '/api/sessions/' + created['session_id']
        headers = {'X-Venue-Token': created['session_token']}
        analyzed = client.post(base+'/analyze', headers=headers, json={'venue': {
            'active_image_height_m':1.8, 'farthest_viewer_distance_m':10.8, 'measurement_basis':'MEASURED'}})
        assert analyzed.status_code == 200, analyzed.text
        summary = analyzed.json()['analysis']['summary']
        assert summary['unsupported_visible_content_elements'] == (1 if case == 'direct' else 0)
        plan = client.post(base+'/plan', headers=headers, json={'constraints': {
            'max_mutations':2, 'max_scale_factor':1.6, 'minimum_target_coverage':.9,
            'protect_titles':True, 'semantic_rewrite':False}})
        assert plan.status_code == 200 and plan.json()['state'] == 'AWAITING_APPROVAL', plan.text
        planned = plan.json()
        done = client.post(base+'/approve-and-apply', headers=headers, json={
            'candidate_ids':planned['decision']['selected_candidate_ids'], 'plan_hash':planned['plan_hash'],
            'idempotency_key':'final01-control', 'approve':True})
        assert done.status_code == 200, done.text
        assert done.json()['state'] == ('REVIEW_REQUIRED' if case == 'direct' else 'VERIFIED'), done.text
        output = client.get(base+'/output', headers=headers)
        assert output.status_code == 200
        assert hashlib.sha256(output.content).hexdigest() == done.json()['verification']['output_sha256']
        assert hashlib.sha256(source.read_bytes()).hexdigest() == original
        # Completed review or verified outputs retain the RC5 immutability contract.
        assert client.post(base+'/plan', headers=headers, json={}).status_code == 409
        assert client.delete(base, headers=headers).status_code == 200
        assert client.get(base+'/output', headers=headers).status_code == 404

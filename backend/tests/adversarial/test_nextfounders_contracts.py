"""Next Founders regression gates for OOXML semantics and public responsiveness."""
from __future__ import annotations

import asyncio
import hashlib
import time
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from zipfile import ZipFile, ZIP_DEFLATED

import httpx
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app import main
from app.domain.models import (
    AnalysisCapability, AnalysisResult, AnalysisSummary, ComparisonState,
    ElementResult, MappingConfidence, MutationCapability, TextElement, VenueProfile,
    WorkflowState,
)
from app.pptx.mutations import ScaleMutation, apply_scale_plan
from app.pptx.package_integrity import package_integrity_reasons
from app.pptx.renderer import _renderer_environment
from app.repair import verifier
from app.storage.sessions import SessionStore


def _deck(path: Path) -> tuple[int, int]:
    prs = Presentation()
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = f'Training example {i + 1}'
        run.font.size = Pt(12)
    prs.save(path)
    return 0, prs.slides[0].shapes[0].shape_id


def _tamper_relationships(source: Path, output: Path) -> None:
    name = 'ppt/_rels/presentation.xml.rels'
    with ZipFile(source) as original, ZipFile(output, 'w', ZIP_DEFLATED) as result:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename == name:
                assert b'slides/slide2.xml' in data
                data = data.replace(b'slides/slide2.xml', b'slides/slide3.xml', 1)
            result.writestr(entry, data)


def test_normal_font_edit_preserves_protected_ooxml(tmp_path: Path):
    source, output = tmp_path/'source.pptx', tmp_path/'updated.pptx'
    slide, shape = _deck(source)
    apply_scale_plan(source, output, [ScaleMutation(slide, shape, 1.2)])
    assert package_integrity_reasons(source, output) == []


def test_retargeted_navigation_is_rejected_even_with_identical_visible_text(tmp_path: Path, monkeypatch):
    source, repaired, tampered = tmp_path/'source.pptx', tmp_path/'repaired.pptx', tmp_path/'tampered.pptx'
    slide, shape = _deck(source)
    apply_scale_plan(source, repaired, [ScaleMutation(slide, shape, 1.2)])
    _tamper_relationships(repaired, tampered)
    assert 'protected_package_relationships_changed' in package_integrity_reasons(source, tampered)
    assert not package_integrity_reasons(source, repaired)

    elem = TextElement(element_id='s1-sh1', slide_index=0, shape_id=shape, source_text='Training example 1',
                       normalized_text='Training example 1', role='BODY', rendered_height_percent=1,
                       analysis_capability=AnalysisCapability.ANALYZABLE,
                       mutation_capability=MutationCapability.SAFE_MUTATION,
                       mapping_confidence=MappingConfidence.EXACT)
    venue = VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8)
    before = AnalysisResult(source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                            reference_profile_id='test', venue=venue,
                            results=[ElementResult(element=elem, state=ComparisonState.BELOW_TARGET,
                                                   required_percent=2, actual_percent=1)],
                            summary=AnalysisSummary(total_detected_text_elements=1, analyzed_elements=1,
                                                    not_analyzed_elements=0, below_target_elements=1,
                                                    meets_target_elements=0, outside_reference_elements=0))
    after = before.model_copy(deep=True)
    after.results[0].actual_percent = 3
    after.results[0].state = ComparisonState.MEETS_TARGET
    after.summary.below_target_elements = 0
    after.summary.meets_target_elements = 1
    monkeypatch.setattr(verifier, '_rendered_text', lambda *args: 'Training example 1')
    monkeypatch.setattr(verifier, 'audit_analysis_transition', lambda *args: SimpleNamespace(
        analysis_universe_preserved=True, layout_safe=True, reasons=[]))
    analyzer = SimpleNamespace(analyze=lambda *args: after)
    report = verifier.verify_output(source, tampered, before, ['s1-sh1'], analyzer, tmp_path/'work')
    assert report.final_state == WorkflowState.REVIEW_REQUIRED
    assert 'protected_package_relationships_changed' in report.reasons
    assert report.protected_text_unchanged is False


def test_renderer_subprocess_environment_excludes_app_secrets(monkeypatch):
    monkeypatch.setenv('GEMINI_API_KEY', 'secret-do-not-inherit')
    monkeypatch.setenv('VENUE_CHILD_PID_FILE', 'secret-test-env')
    monkeypatch.setenv('HOME', '/tmp/test-home')
    env = _renderer_environment()
    assert env['HOME'] == '/tmp/test-home'
    assert 'GEMINI_API_KEY' not in env
    assert 'VENUE_CHILD_PID_FILE' not in env


def test_expensive_upload_validation_does_not_block_health(monkeypatch, tmp_path):
    # Isolation: replace only expensive intake validation, never mock HTTP scheduling.
    store = SessionStore(tmp_path/'sessions', max_sessions=3)
    def slow_validation(_path):
        time.sleep(0.4)
        return SimpleNamespace(has_external_relationships=False)
    monkeypatch.setattr(main, 'STORE', store)
    monkeypatch.setattr(main, 'validate_pptx', slow_validation)
    async def run():
        transport = httpx.ASGITransport(app=main.app)
        async with httpx.AsyncClient(transport=transport, base_url='http://testserver') as client:
            uploading = asyncio.create_task(client.post('/api/upload',
                                    files={'file': ('training.pptx', b'placeholder')}))
            await asyncio.sleep(.06)
            t = time.monotonic()
            response = await client.get('/health')
            elapsed = time.monotonic()-t
            uploaded = await uploading
            assert response.status_code == 200
            assert uploaded.status_code == 200, uploaded.text
            assert elapsed < .25, f'health blocked by upload for {elapsed:.3f}s'
    try:
        asyncio.run(run())
    finally:
        store.stop_reaper()

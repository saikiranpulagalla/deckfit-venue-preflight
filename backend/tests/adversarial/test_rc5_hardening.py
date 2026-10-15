"""RC5 regression gates: parser complexity, provenance, immutable result, concurrency."""
from __future__ import annotations

import hashlib
import os
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches, Pt

from app import main
from app.agent.planner import GeminiPlanner
from app.pptx.mutations import ScaleMutation, apply_scale_plan
from app.pptx.package_integrity import package_integrity_reasons
from app.security.intake import IntakeError, _external_svg_url, validate_pptx
from app.storage.sessions import SessionStore


def base_pptx(path: Path, *, two_slides: bool = False) -> tuple[Path, list[int]]:
    prs = Presentation()
    identifiers = []
    for index in range(2 if two_slides else 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = f'Workshop body {index + 1}'
        run.font.size = Pt(10)
        identifiers.append(shape.shape_id)
    prs.save(path)
    return path, identifiers


def rewrite(src: Path, dest: Path, modifications: dict[str, bytes], *, compression=zipfile.ZIP_DEFLATED):
    with zipfile.ZipFile(src) as zf, zipfile.ZipFile(dest, 'w', compression=compression) as out:
        for member in zf.infolist():
            out.writestr(member.filename, modifications.get(member.filename, zf.read(member.filename)))
        for name, content in modifications.items():
            if name not in zf.namelist():
                out.writestr(name, content)


def embed_svg(src: Path, dest: Path, content: bytes):
    with zipfile.ZipFile(src) as zf:
        root = ET.fromstring(zf.read('[Content_Types].xml'))
    ns = '{http://schemas.openxmlformats.org/package/2006/content-types}'
    ET.SubElement(root, ns + 'Default', {'Extension': 'svg', 'ContentType': 'image/svg+xml'})
    rewrite(src, dest, {'ppt/media/rc5.svg': content,
                        '[Content_Types].xml': ET.tostring(root, encoding='utf-8')})


def test_embedded_svg_complexity_bounded(tmp_path):
    src, _ = base_pptx(tmp_path/'source.pptx')
    p = tmp_path/'whitespace.pptx'
    embed_svg(src, p, ("<svg xmlns='http://www.w3.org/2000/svg'><rect style='fill:url(" +
                       " " * 20000 + "INVALID)'/></svg>").encode())
    start = time.monotonic()
    assert validate_pptx(p).slide_count == 1
    assert time.monotonic() - start < 1.5
    assert _external_svg_url('url(  \"  https://example.invalid/media)')
    assert _external_svg_url('url(   //external.invalid)')
    assert not _external_svg_url('url(  NOT_EXTERNAL)')


def test_embedded_svg_entity_expansion_is_refused(tmp_path):
    src, _ = base_pptx(tmp_path/'source.pptx')
    p = tmp_path/'dtd.pptx'
    embed_svg(src, p, b'<!DOCTYPE svg [<!ENTITY a "TEST">]><svg xmlns="http://www.w3.org/2000/svg" style="&a;"/>')
    with pytest.raises(IntakeError):
        validate_pptx(p)


@pytest.mark.parametrize('compression', [zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA])
def test_unsupported_zip_codec_rejected_during_intake(tmp_path, compression):
    src, _ = base_pptx(tmp_path/'source.pptx')
    p = tmp_path/'unsupported.pptx'
    rewrite(src, p, {}, compression=compression)
    with pytest.raises(IntakeError, match='compression'):
        validate_pptx(p)


def test_deep_slide_xml_rejected_at_intake(tmp_path):
    src, _ = base_pptx(tmp_path/'source.pptx')
    with zipfile.ZipFile(src) as zf:
        data = zf.read('ppt/slides/slide1.xml')
    nested = b'<w>' * 250 + data + b'</w>' * 250
    p = tmp_path/'deep.pptx'
    rewrite(src, p, {'ppt/slides/slide1.xml': nested})
    with pytest.raises(IntakeError, match='nesting'):
        validate_pptx(p)


def test_reordered_slides_are_authorized_by_logical_position(tmp_path):
    src, ids = base_pptx(tmp_path/'source.pptx', two_slides=True)
    with zipfile.ZipFile(src) as zf:
        root = ET.fromstring(zf.read('ppt/presentation.xml'))
    ns = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
    listing = root.find(ns + 'sldIdLst')
    first = listing[0]
    listing.remove(first)
    listing.append(first)
    reordered = tmp_path/'reordered.pptx'
    rewrite(src, reordered, {'ppt/presentation.xml': ET.tostring(root, encoding='utf-8')})
    assert validate_pptx(reordered).slide_count == 2
    physical_id = ids[1]
    edited = tmp_path/'repaired.pptx'
    apply_scale_plan(reordered, edited, [ScaleMutation(0, physical_id, 1.25)])
    assert package_integrity_reasons(reordered, edited, {(0, physical_id)}, {(0, physical_id):1.25}) == []
    # Unauthorized modifications to the other slide must still be rejected.
    unauthorized = tmp_path/'tampered.pptx'
    apply_scale_plan(reordered, unauthorized, [ScaleMutation(1, ids[0], 1.25)])
    assert package_integrity_reasons(reordered, unauthorized, {(0, physical_id)}, {(0, physical_id):1.25})


def test_circuit_breaker_older_success_cannot_erase_new_failure():
    planner = GeminiPlanner()
    generation = planner._circuit_generation
    planner._open_circuit()
    assert planner._circuit_is_open()
    planner._close_circuit(generation)
    assert planner._circuit_is_open()
    planner._close_circuit(planner._circuit_generation)
    assert not planner._circuit_is_open()


def test_upload_reservation_keeps_capacity_from_demo(tmp_path):
    store = SessionStore(tmp_path / 'sessions', max_sessions=2, per_owner_sessions=2)
    try:
        assert store.begin_intake('upload_a') is None
        assert store.begin_intake('upload_b') is None
        with pytest.raises(RuntimeError, match='session_capacity_reached'):
            store.create('demo.pptx', b'demo', 'h', owner_key='demo')
        a = store.create('a.pptx', b'a', 'ha', owner_key='upload_a', reserved_intake=True)
        store.finish_intake('upload_a')
        with pytest.raises(RuntimeError, match='session_capacity_reached'):
            store.create('demo.pptx', b'demo', 'h', owner_key='demo')
        b = store.create('b.pptx', b'b', 'hb', owner_key='upload_b', reserved_intake=True)
        store.finish_intake('upload_b')
        assert a.source_path.is_file() and b.source_path.is_file()
        assert len(store._items) == 2
    finally:
        store.stop_reaper()


@pytest.mark.resource
def test_completed_verified_output_immutable_to_replan_and_reanalysis(tmp_path, monkeypatch):
    store = SessionStore(tmp_path / 'sessions')
    monkeypatch.setattr(main, 'STORE', store)
    venue = {'active_image_height_m':1.8,'farthest_viewer_distance_m':10.8,'measurement_basis':'MEASURED'}
    constraints = {'max_mutations':2,'max_scale_factor':1.6,'minimum_target_coverage':0.75,
                   'protect_titles':True,'semantic_rewrite':False}
    try:
        with TestClient(main.app) as client:
            created = client.post('/api/demo').json()
            sid = created['session_id']
            h = {'X-Venue-Token':created['session_token']}
            base = f'/api/sessions/{sid}'
            assert client.post(base+'/analyze', headers=h, json={'venue':venue}).status_code == 200
            plan = client.post(base+'/plan', headers=h, json={'constraints':constraints}).json()
            assert plan['state'] == 'AWAITING_APPROVAL', plan
            done = client.post(base+'/approve-and-apply', headers=h, json={
                'approve':True,'plan_hash':plan['plan_hash'],
                'candidate_ids':plan['decision']['selected_candidate_ids'],
                'idempotency_key':'rc5-immutable-approved-result'})
            assert done.status_code == 200 and done.json()['state'] == 'VERIFIED', done.text
            downloaded = client.get(base+'/output', headers=h)
            assert downloaded.status_code == 200
            digest = hashlib.sha256(downloaded.content).hexdigest()
            assert client.post(base+'/plan', headers=h, json={'constraints':constraints}).status_code == 409
            assert client.post(base+'/analyze', headers=h, json={'venue':venue}).status_code == 409
            assert hashlib.sha256(client.get(base+'/output', headers=h).content).hexdigest() == digest
    finally:
        store.stop_reaper()


def test_deployed_source_identity_cannot_be_established_by_label_only(tmp_path, monkeypatch):
    """Historical RC3 carrying RC5's release label must fail independently-bound verification."""
    import json
    import sys
    from app.security.build_identity import source_fingerprint
    sys.path.insert(0, str(Path(__file__).resolve().parents[3] / 'scripts'))
    import live_deployment_check as live
    root = Path(__file__).resolve().parents[3]
    expected = source_fingerprint(root)
    assert len(expected) == 64
    good_headers = {
        'referrer-policy':'no-referrer', 'x-content-type-options':'nosniff',
        'x-frame-options':'DENY',
        'content-security-policy':"frame-ancestors 'none'; object-src 'none'",
        'cache-control':'no-store'
    }
    def faux_request(base_url, path, **kwargs):
        assert path == '/health', 'Must reject spoofed release label before probing anything else'
        body = json.dumps({'release_id':'deckfit-nextfounders-rc5-local','planner_mode':'deterministic_fallback'}).encode()
        return live.HttpResult(200, body, good_headers, 1.0)
    monkeypatch.setattr(live, '_request', faux_request)
    with pytest.raises(RuntimeError, match='deployment_source_fingerprint_mismatch'):
        live.run_check('http://old-code.invalid', expected_release_id='deckfit-nextfounders-rc5-local',
                       expected_source_fingerprint=expected, require_gemini=False, full=False)


def test_verified_artifact_stays_downloadable_through_refused_replan(tmp_path, monkeypatch):
    """Fast API-level state gate independent of LibreOffice execution time."""
    from app.domain.models import WorkflowState
    store = SessionStore(tmp_path/'sessions')
    monkeypatch.setattr(main,'STORE',store)
    try:
        with TestClient(main.app) as client:
            s = store.create('source.pptx', b'not-a-real-pptx', 'seed', owner_key='test')
            token = s.take_issued_capability_token()
            saved = s.root/'output.pptx'
            saved.write_bytes(b'the-verified-artifact')
            s.output_path = saved
            s.output_sha256 = hashlib.sha256(saved.read_bytes()).hexdigest()
            s.state = WorkflowState.VERIFIED
            uri=f'/api/sessions/{s.session_id}'
            headers={'X-Venue-Token':token}
            assert client.post(uri+'/plan', headers=headers, json={}).status_code == 409
            assert client.post(uri+'/analyze', headers=headers, json={
                'venue':{'active_image_height_m':1.8,'farthest_viewer_distance_m':10.8,
                         'measurement_basis':'MEASURED'}}).status_code == 409
            r = client.get(uri+'/output', headers=headers)
            assert r.status_code == 200 and r.content == b'the-verified-artifact'
    finally:
        store.stop_reaper()


def test_session_expiry_uses_elapsed_time_despite_wall_clock_rollback(tmp_path):
    store = SessionStore(tmp_path/'sessions', ttl_seconds=600, provisional_ttl_seconds=60)
    try:
        session = store.create('clock.pptx', b'content', 'digest', owner_key='owner')
        token = session.take_issued_capability_token()
        assert store.authorize(session.session_id, token)
        deadline = session.monotonic_expires_at
        store._clock = lambda: session.created_at - 3600
        store._elapsed_clock = lambda: deadline + 1
        assert store.cleanup_expired() == 1
        assert not session.source_path.exists()
    finally:
        store.stop_reaper()


@pytest.mark.resource
def test_reordered_slide_output_retains_correct_render_order(tmp_path):
    import fitz
    from app.pptx.renderer import render_pptx_to_pdf
    source, ids = base_pptx(tmp_path/'physical.pptx', two_slides=True)
    with zipfile.ZipFile(source) as zf:
        root = ET.fromstring(zf.read('ppt/presentation.xml'))
    slides = root.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst')
    first = slides[0]
    slides.remove(first)
    slides.append(first)
    reordered = tmp_path/'ordered.pptx'
    rewrite(source, reordered, {'ppt/presentation.xml':ET.tostring(root, encoding='utf-8')})
    output = tmp_path/'scaled.pptx'
    apply_scale_plan(reordered, output, [ScaleMutation(0, ids[1], 1.25)])
    assert package_integrity_reasons(reordered, output, {(0,ids[1])}, {(0,ids[1]):1.25}) == []
    pdf = render_pptx_to_pdf(output, tmp_path/'pdf')
    with fitz.open(pdf) as document:
        assert len(document) == 2
        assert 'Workshop body 2' in document[0].get_text()
        assert 'Workshop body 1' in document[1].get_text()


def test_rc5_release_check_refuses_unbound_version_claim():
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[3] / 'scripts'))
    import live_deployment_check as live
    with pytest.raises(RuntimeError, match='rc5_source_fingerprint_required'):
        live.run_check('http://server.invalid', expected_release_id='deckfit-nextfounders-rc5-local',
                       expected_source_fingerprint=None, require_gemini=False, full=False)

"""Frozen R6 regressions: new invalid cases must not pass safety, admission or proof gates."""
from __future__ import annotations

import hashlib
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app import main
from app.domain.models import (
    AnalysisCapability, AnalysisResult, AnalysisSummary, ComparisonState,
    ElementResult, MappingConfidence, MutationCapability, RepairCandidate,
    TextElement, VenueProfile, WorkflowState, PlanDecision,
)
from app.pptx.parser import parse_pptx_shapes
from app.pptx.renderer import RenderError
from app.security.intake import IntakeError, validate_pptx
from app.storage.sessions import SessionBusy, SessionStore

ROOT = Path(__file__).resolve().parents[3]
FIX = ROOT / 'backend/tests/fixtures/r6'
DEMO = ROOT / 'demo/venue-agent-demo.pptx'
VENUE = {'active_image_height_m': 1.8, 'farthest_viewer_distance_m': 10.8,
         'measurement_basis': 'MEASURED'}


def _fake_analysis(src: Path) -> AnalysisResult:
    element = TextElement(
        element_id='s1-sh1', slide_index=0, shape_id=1, source_text='tiny',
        normalized_text='tiny', role='BODY', rendered_height_percent=1,
        analysis_capability=AnalysisCapability.ANALYZABLE,
        mutation_capability=MutationCapability.SAFE_MUTATION,
        mapping_confidence=MappingConfidence.EXACT,
    )
    return AnalysisResult(
        source_sha256=hashlib.sha256(src.read_bytes()).hexdigest(),
        reference_profile_id='dummy', venue=VenueProfile(**VENUE),
        results=[ElementResult(element=element, state=ComparisonState.BELOW_TARGET,
                               required_percent=2, actual_percent=1, deficit_ratio=2)],
        summary=AnalysisSummary(total_detected_text_elements=1, analyzed_elements=1,
                                not_analyzed_elements=0, below_target_elements=1,
                                meets_target_elements=0, outside_reference_elements=0),
    )


def test_painted_empty_text_box_is_inventoried_but_blank_unpainted_is_not():
    records = parse_pptx_shapes(FIX / 'r5-empty-filled-textbox-mask.pptx')
    assert any(r.unsupported_visible_content and 'text_box' in ''.join(r.notes)
               for r in records), records
    clean = parse_pptx_shapes(DEMO)
    assert not any(r.unsupported_visible_content for r in clean)


def test_truly_unpainted_empty_text_box_keeps_normal_demo_supported(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation(DEMO)
    deck.slides[0].shapes.add_textbox(Inches(2), Inches(2), Inches(2), Inches(1))
    target = tmp_path / 'empty-unpainted-box.pptx'
    deck.save(target)
    records = parse_pptx_shapes(target)
    assert len(records) == 3
    assert not any(record.unsupported_visible_content for record in records)


def test_explicitly_hidden_master_graphics_not_inventoried():
    records = parse_pptx_shapes(FIX / 'hidden-master-graphics.pptx')
    assert not any('MASTER_ONLY_SENTINEL' in r.text for r in records)


def test_malformed_referenced_layout_rejected_before_session_allocation():
    with pytest.raises(IntakeError, match='malformed layout/master/theme'):
        validate_pptx(FIX / 'r5-corrupt-layout-accepted.pptx')
    with TestClient(main.app) as client:
        with (FIX / 'r5-corrupt-layout-accepted.pptx').open('rb') as f:
            response = client.post('/api/upload', files={'file': ('corrupt.pptx', f)})
        assert response.status_code == 400, response.text


def test_degraded_renderer_is_non_2xx_for_deployment_probe():
    with patch.object(main, '_deep_renderer_readiness', return_value=(False, 'render_probe_failed:RenderError')):
        with patch.object(main, 'libreoffice_version', return_value='mock LibreOffice'):
            with TestClient(main.app) as c:
                r = c.get('/health/deep-readiness')
                assert r.status_code == 503, r.text
                assert r.json()['status'] == 'degraded'
    with patch.object(main, '_deep_renderer_readiness', return_value=(True, 'real_demo_render_pass')):
        with patch.object(main, 'libreoffice_version', return_value='mock LibreOffice'):
            with TestClient(main.app) as c:
                assert c.get('/health/deep-readiness').status_code == 200


def test_early_intake_rejects_before_expensive_upload_validation(tmp_path):
    store = SessionStore(tmp_path / 'sessions', max_sessions=4, per_owner_sessions=1)
    with patch.object(main, 'STORE', store):
        with TestClient(main.app) as client:
            assert client.post('/api/demo').status_code == 200
            with patch.object(main, 'validate_pptx', side_effect=AssertionError('intake must not run')):
                response = client.post('/api/upload', files={'file': ('deck.pptx', DEMO.read_bytes())})
            assert response.status_code == 429, response.text
            assert response.headers.get('retry-after') == '3'
    store.stop_reaper()


def test_early_intake_other_owner_not_blocked_and_atomic_reservations(tmp_path):
    store = SessionStore(tmp_path / 'sessions', max_sessions=3, per_owner_sessions=2,
                         max_parallel_intakes_per_owner=1)
    assert store.begin_intake('alice') is None
    assert store.begin_intake('alice') == 'upload_admission_busy'
    assert store.begin_intake('bob') is None
    store.finish_intake('alice')
    store.finish_intake('bob')
    assert not store._pending_intakes


def test_failed_intake_does_not_reserve_upload_capacity(tmp_path):
    store = SessionStore(tmp_path / 'sessions', max_sessions=3, per_owner_sessions=1)
    with patch.object(main, 'STORE', store):
        with TestClient(main.app) as client:
            bad = client.post('/api/upload', files={'file': ('bad.pptx', b'garbage')})
            assert bad.status_code == 400
            ok = client.post('/api/upload', files={'file': ('ok.pptx', DEMO.read_bytes())})
            assert ok.status_code == 200, ok.text
    store.stop_reaper()


def test_lock_queue_bounded_and_timed_waiters_release_leases(tmp_path):
    store = SessionStore(tmp_path / 'sessions', max_waiters_per_session=2, lock_wait_timeout_seconds=0.12)
    session = store.create('x', b'x', '0'*64)
    entered, release = threading.Event(), threading.Event()

    def holder():
        with store.locked(session.session_id):
            entered.set()
            release.wait(2)

    t = threading.Thread(target=holder)
    t.start()
    assert entered.wait(1)
    outcomes = []

    def contender():
        try:
            with store.locked(session.session_id):
                return 'acquired'
        except SessionBusy:
            return 'busy'

    with ThreadPoolExecutor(max_workers=8) as pool:
        futures = [pool.submit(contender) for _ in range(8)]
        outcomes = [future.result(timeout=2) for future in futures]
    release.set()
    t.join(timeout=2)
    assert outcomes.count('busy') == 8
    assert session.active_operations == 0
    assert store.delete(session.session_id)


def test_public_api_reports_session_lock_overload_as_retryable_http(tmp_path):
    store = SessionStore(tmp_path / 'sessions', max_waiters_per_session=0,
                         lock_wait_timeout_seconds=0.1)
    entered, release = threading.Event(), threading.Event()
    with patch.object(main, 'STORE', store):
        with TestClient(main.app) as c:
            created = c.post('/api/demo').json()
            sid, auth = created['session_id'], {'X-Venue-Token': created['session_token']}
            def hold():
                with store.locked(sid):
                    entered.set()
                    release.wait(2)
            thread = threading.Thread(target=hold)
            thread.start()
            assert entered.wait(1)
            response = c.get(f'/api/sessions/{sid}', headers=auth)
            assert response.status_code == 503
            assert response.headers.get('retry-after') == '3'
            release.set()
            thread.join(timeout=2)
            assert c.get(f'/api/sessions/{sid}', headers=auth).status_code == 200
            assert c.delete(f'/api/sessions/{sid}', headers=auth).status_code == 200
    store.stop_reaper()


def test_analyze_simulation_busy_is_retryable_and_does_not_publish_partial_candidates(tmp_path):
    store = SessionStore(tmp_path / 'sessions')
    with patch.object(main, 'STORE', store):
        with TestClient(main.app) as c:
            created = c.post('/api/demo').json()
            sid, auth = created['session_id'], {'X-Venue-Token': created['session_token']}
            candidate = RepairCandidate(candidate_id='c1', issue_element_id='s1-sh1', operation='SCALE_TEXT', parameters={'scale': 1.1})
            with patch.object(main.ANALYZER, 'analyze', side_effect=lambda src,*args,**kwargs: _fake_analysis(src)):
                with patch.object(main, 'generate_candidates', return_value=[candidate]):
                    with patch.object(main, 'simulate_candidate', side_effect=RenderError('Renderer capacity is busy; retry later')):
                        response = c.post(f'/api/sessions/{sid}/analyze', headers=auth, json={'venue': VENUE})
            assert response.status_code == 503, response.text
            assert response.headers.get('retry-after') == '3'
            with store.locked(sid) as s:
                assert s.state == WorkflowState.UPLOADED
                assert s.analysis is None and not s.candidates and s.plan_hash is None
            c.delete(f'/api/sessions/{sid}', headers=auth)
    store.stop_reaper()


@pytest.mark.parametrize('fail_index', [0, 1, 2])
def test_candidate_simulation_exception_is_transactional_at_every_index(fail_index, tmp_path):
    store = SessionStore(tmp_path / 'sessions')
    with patch.object(main, 'STORE', store):
        with TestClient(main.app) as c:
            created = c.post('/api/demo').json()
            sid, auth = created['session_id'], {'X-Venue-Token': created['session_token']}
            candidates = [RepairCandidate(candidate_id=f'c{i}', issue_element_id='s1-sh1',
                                          operation='SCALE_TEXT', parameters={'scale': 1.1})
                          for i in range(3)]
            called = [0]
            def simulated(_source, _analysis, candidate, _analyzer):
                called[0] += 1
                if called[0] == fail_index + 1:
                    raise RenderError('Renderer capacity is busy; retry later')
                return candidate
            with patch.object(main.ANALYZER, 'analyze', side_effect=lambda src,*a,**kw: _fake_analysis(src)):
                with patch.object(main, 'generate_candidates', return_value=candidates):
                    with patch.object(main, 'simulate_candidate', side_effect=simulated):
                        r = c.post(f'/api/sessions/{sid}/analyze', headers=auth, json={'venue': VENUE})
                    assert r.status_code == 503, r.text
                    with store.locked(sid) as s:
                        assert s.state == WorkflowState.UPLOADED
                        assert s.analysis is None and s.candidates == []
                    # Same-session retry after temporary contention clears.
                    with patch.object(main, 'simulate_candidate', side_effect=lambda _s,_a,c,_v: c):
                        recovered = c.post(f'/api/sessions/{sid}/analyze', headers=auth, json={'venue': VENUE})
                    assert recovered.status_code == 200, recovered.text
                    assert recovered.json()['state'] == 'ANALYZED'
                    assert len(recovered.json()['candidates']) == 3
            assert c.delete(f'/api/sessions/{sid}', headers=auth).status_code == 200
    store.stop_reaper()


def test_failed_replan_clears_approval_and_returns_to_analyzed(tmp_path):
    store = SessionStore(tmp_path / 'sessions')
    with patch.object(main, 'STORE', store):
        with TestClient(main.app) as client:
            created = client.post('/api/demo').json()
            sid, auth = created['session_id'], {'X-Venue-Token': created['session_token']}
            with store.locked(sid) as s:
                s.analysis = _fake_analysis(s.source_path)
                s.candidates = [RepairCandidate(candidate_id='c1', issue_element_id='s1-sh1', operation='SCALE_TEXT', parameters={'scale': 1.1}, safe=True, simulation_fingerprint='fp')]
                s.plan_hash = 'old-plan'
                s.state = WorkflowState.AWAITING_APPROVAL
            decision = PlanDecision(selected_candidate_ids=['c1'], rationale='bounded', requires_human_review=False, unresolved_issue_ids=[], projected_target_coverage=1.0)
            with patch.object(main.PLANNER, 'choose', return_value=decision):
                with patch.object(main, 'simulate_plan', side_effect=RenderError('Renderer capacity is busy; retry later')):
                    response = client.post(f'/api/sessions/{sid}/plan', headers=auth, json={})
            assert response.status_code == 503, response.text
            with store.locked(sid) as s:
                assert s.state == WorkflowState.ANALYZED
                assert s.plan_hash is None and not s.selected_candidate_ids and s.approved_plan_hash is None
            assert client.delete(f'/api/sessions/{sid}', headers=auth).status_code == 200
    store.stop_reaper()


def test_large_approval_identifier_inputs_rejected_at_schema():
    with pytest.raises(ValidationError):
        main.ApproveRequest(candidate_ids=['c1'], plan_hash='p'*1048576, idempotency_key='k'*1048576)
    with pytest.raises(ValidationError):
        main.ApproveRequest(candidate_ids=['c1'], plan_hash='a'*64, idempotency_key='bad value')
    ok = main.ApproveRequest(candidate_id='c1', plan_hash='a'*64, idempotency_key='test-key.1')
    assert ok.candidate_ids == ['c1']


@pytest.mark.parametrize('failed_part', ['source.pptx', '.venue-agent-session-v1'])
def test_failed_create_write_rolls_back_owned_root(tmp_path, failed_part):
    store = SessionStore(tmp_path / 'sessions', max_sessions=2)
    from app.storage import sessions
    original = sessions._secure_write
    def write_with_error(path, content):
        if path.name == failed_part:
            raise OSError('synthetic disk error')
        return original(path, content)
    with patch.object(sessions, '_secure_write', side_effect=write_with_error):
        with pytest.raises(OSError):
            store.create('x', b'x', 'a'*64)
    assert store._items == {}
    assert list(store.base.glob('session-*')) == []
    good = store.create('x', b'x', 'a'*64)
    assert store.delete(good.session_id)


@pytest.mark.resource
def test_actual_rendered_mask_does_not_verify_in_public_api():
    with TestClient(main.app) as client:
        p = FIX / 'r5-empty-filled-textbox-mask.pptx'
        with p.open('rb') as f:
            created = client.post('/api/upload', files={'file': (p.name, f)})
        assert created.status_code == 200, created.text
        sid, auth = created.json()['session_id'], {'X-Venue-Token': created.json()['session_token']}
        try:
            analyzed = client.post(f'/api/sessions/{sid}/analyze', headers=auth, json={'venue': VENUE})
            assert analyzed.status_code == 200, analyzed.text
            assert analyzed.json()['analysis']['summary']['unsupported_visible_content_elements'] > 0
            if analyzed.json()['state'] == 'ANALYZED':
                planned = client.post(f'/api/sessions/{sid}/plan', headers=auth,
                                      json={'constraints': {'max_mutations': 2, 'max_scale_factor': 1.6,
                                                           'minimum_target_coverage': .75}})
                assert planned.status_code == 200, planned.text
                if planned.json()['state'] == 'AWAITING_APPROVAL':
                    req = {'candidate_ids': planned.json()['decision']['selected_candidate_ids'],
                           'plan_hash': planned.json()['plan_hash'], 'idempotency_key': 'r6-test'}
                    result = client.post(f'/api/sessions/{sid}/approve-and-apply', headers=auth, json=req)
                    assert result.status_code == 200, result.text
                    assert result.json()['state'] != 'VERIFIED'
        finally:
            client.delete(f'/api/sessions/{sid}', headers=auth)

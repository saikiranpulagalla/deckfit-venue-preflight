from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE

ROOT=Path(__file__).resolve().parents[1]
out=ROOT/'demo'/'venue-agent-demo.pptx'
out.parent.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
slide=prs.slides.add_slide(prs.slide_layouts[6])

def box(text,left,top,width,height,size):
    s=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height))
    s.text_frame.auto_size=MSO_AUTO_SIZE.NONE
    r=s.text_frame.paragraphs[0].add_run(); r.text=text; r.font.size=Pt(size); r.font.name="DejaVu Sans"; return s
box('Venue Agent: Room-Aware Repair',0.75,0.55,11.8,0.8,34)
box('This primary body text is intentionally too small for the modeled farthest-viewer reference.',0.85,2.0,11.5,1.35,13)
box('A second body region demonstrates a bounded multi-issue repair plan.',0.85,4.25,10.8,1.1,13)
prs.save(out)
print(out)
